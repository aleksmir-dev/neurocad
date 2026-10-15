# neurocad/core/engine/lib/botinv/importer.py

"""
Parse an inventory Excel/CSV file.

Fixed logical schema: inv_num, name, room.
Column order in the file does NOT matter — columns are located
by header name in the first row.

Returns a list of dicts ready for BotInvItem(**item):
    {
        "inv_num": str,        # raw, as it appeared in the file
        "inv_num_norm": str,   # inv_num.upper() — for lookups
        "name": str,
        "room_before": str | None,
    }

Case handling
-------------
Two rows with inv_num "ЧФ 0123" and "чф 0123" are the SAME item
for the inventory — the search is case-insensitive. This module
therefore collapses them into one row (first occurrence wins),
matching the UNIQUE constraint in the DB model
(uq_botinv_sheet_invnum_norm).

Whitespace INSIDE the number is preserved: "ЧФ 0123" and
"ЧФ0123" remain distinct items. Only the case is normalized.

Raises ValueError with a human-readable Russian message on any
structural problem (missing headers, unreadable file, empty file).
"""

import csv
import io
from typing import Optional

from openpyxl import load_workbook


REQUIRED_HEADERS = ("inv_num", "name", "room")


def parse_inventory(data: bytes, filename: str) -> list[dict]:
    """
    Parse an inventory file into a list of item dicts.

    Supports .xlsx (openpyxl) and .csv (stdlib csv).
    Detects the format by file extension.

    Raises ValueError on any structural problem.
    """
    name_lower = (filename or "").lower()

    if name_lower.endswith(".xlsx"):
        rows = _read_xlsx(data)
    elif name_lower.endswith(".csv"):
        rows = _read_csv(data)
    else:
        raise ValueError(
            "Поддерживаются только файлы .xlsx и .csv. "
            "Сохраните ведомость в одном из этих форматов."
        )

    return _rows_to_items(rows)


def _read_xlsx(data: bytes) -> list[list]:
    """Read the first worksheet of an .xlsx file into a list of rows."""
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:
        raise ValueError(f"Не удалось открыть .xlsx: {exc}") from exc

    ws = wb.active
    if ws is None:
        raise ValueError("В файле .xlsx нет активного листа.")

    rows = []
    for row in ws.iter_rows(values_only=True):
        rows.append(list(row))
    wb.close()
    return rows


def _read_csv(data: bytes) -> list[list]:
    """
    Read a .csv file into a list of rows.

    Tries utf-8-sig first (handles BOM from Excel), then cp1251
    (common for Russian Excel exports).
    """
    text: Optional[str] = None
    for enc in ("utf-8-sig", "cp1251"):
        try:
            text = data.decode(enc)
            break
        except UnicodeDecodeError:
            continue

    if text is None:
        raise ValueError(
            "Не удалось определить кодировку файла. "
            "Сохраните CSV в UTF-8 или Windows-1251."
        )

    # Sniff delimiter: Excel in RU locale usually uses ";",
    # but "," is also common.
    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=";,")
        delimiter = dialect.delimiter
    except csv.Error:
        delimiter = ";"

    reader = csv.reader(io.StringIO(text), delimiter=delimiter)
    return [row for row in reader]


def _rows_to_items(rows: list[list]) -> list[dict]:
    """
    Validate headers, map columns, deduplicate by inv_num_norm.

    Two rows whose normalized numbers are equal ("ЧФ 0123" and
    "чф 0123") collapse into one — the first wins. Whitespace
    inside the number is preserved, so "ЧФ 0123" and "ЧФ0123"
    stay distinct.
    """
    if not rows:
        raise ValueError("Файл пустой — нет ни одной строки.")

    header_row = rows[0]
    header_map = _build_header_map(header_row)

    items: list[dict] = []
    seen: set[str] = set()

    for raw in rows[1:]:
        item = _row_to_item(raw, header_map)
        if item is None:
            continue
        key = item["inv_num_norm"]
        if key in seen:
            continue
        seen.add(key)
        items.append(item)

    if not items:
        raise ValueError(
            "В файле нет ни одной корректной строки с данными. "
            "Проверьте, что заполнены колонки inv_num, name, room."
        )

    return items


def _build_header_map(header_row: list) -> dict:
    """Return {logical_name: column_index} for the required headers."""
    normalized = []
    for cell in header_row:
        if cell is None:
            normalized.append("")
        else:
            normalized.append(str(cell).strip().lower())

    header_map: dict = {}
    for logical in REQUIRED_HEADERS:
        try:
            idx = normalized.index(logical)
        except ValueError:
            raise ValueError(
                "В первой строке не найдены обязательные столбцы: "
                f"inv_num, name, room. "
                f"Ожидались заголовки {list(REQUIRED_HEADERS)}, "
                f"получены: {[h for h in normalized if h]}."
            )
        header_map[logical] = idx

    return header_map


def _row_to_item(raw: list, header_map: dict) -> Optional[dict]:
    """
    Map one raw row to an item dict. Returns None to skip the row.

    inv_num is stored as-is (leading/trailing whitespace stripped,
    case preserved). inv_num_norm is inv_num.upper().
    """
    if _is_blank_row(raw):
        return None

    inv_num = _cell(raw, header_map["inv_num"])
    if not inv_num:
        return None

    name = _cell(raw, header_map["name"]) or ""
    room = _cell(raw, header_map["room"]) or None

    return {
        "inv_num": inv_num,
        "inv_num_norm": inv_num.upper(),
        "name": name,
        "room_before": room,
    }


def _is_blank_row(raw: list) -> bool:
    """A row is blank if all its cells are empty or None."""
    return all(_cell_str(c) == "" for c in raw)


def _cell(raw: list, idx: int) -> Optional[str]:
    """Return a stripped string for cell at idx, or None."""
    if idx >= len(raw):
        return None
    value = raw[idx]
    if value is None:
        return None
    s = str(value).strip()
    return s if s else None


def _cell_str(value) -> str:
    if value is None:
        return ""
    return str(value).strip()